#!/usr/bin/env python3
"""
jkt-anthropic-claim : Anthropic 구독료 월별 ใบเบิก 정산 묶음을 만든다.

  인보이스/영수증 PDF 파싱 -> ใบเบิก 엑셀(템플릿 복사본) 작성 + 서명 삽입
  -> LibreOffice 로 PDF 변환 -> ใบเบิก/Invoice/Receipt/서명PR/(문자) 합본 -> 검증

원본(템플릿·inbox 서류)은 읽기만 한다. 결과물은 전부 out/ 에 새로 만든다.

사용 예
  python3 build_claim.py --month 2610 --thb 6234.50 --pr inbox/2610/PR-26102201.pdf
  python3 build_claim.py --month 2610                      # THB 미정 -> F10 비움, _DRAFT
  python3 build_claim.py --month 2610 --thb 6234.50 --sms inbox/2610/sms.png \
      --mask 40,610,700,760 --mask 40,980,700,1100         # 개인 결제 말풍선 가리기(px)
"""
import argparse, copy, datetime as dt, glob, io, os, re, shutil, subprocess, sys, tempfile, zipfile
import xml.etree.ElementTree as ET

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment
from openpyxl.utils.units import pixels_to_EMU
from PIL import Image
from pypdf import PdfReader, PdfWriter

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(HERE, "..", "..", "..", ".."))
TEMPLATE_NAME = "แบบฟอร์มใบเบิก.xlsx"
NAME = "JEONGHO CHIN"
MASK_RGB = (198, 202, 206)
SIG_H_PX = 52  # 서명 높이(px), 비율 유지
MONTHS = {m: i for i, m in enumerate(
    ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"], 1)}
XDR = "{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def die(msg):
    sys.exit(f"[중단] {msg}")


# ---------------------------------------------------------------- PDF 파싱
def pdf_text(path):
    """pdftotext -layout 우선 (하이픈/대시 보존). 없으면 pdfplumber."""
    if shutil.which("pdftotext"):
        r = subprocess.run(["pdftotext", "-layout", path, "-"], capture_output=True, text=True)
        if r.returncode == 0 and r.stdout.strip():
            return r.stdout
    import pdfplumber
    with pdfplumber.open(path) as pdf:
        return "\n".join(p.extract_text() or "" for p in pdf.pages)


def money(s):
    return float(s.replace(",", ""))


def long_date(s):
    return dt.datetime.strptime(re.sub(r"\s+", " ", s.strip()), "%B %d, %Y").date()


def parse_period(text, after):
    """'Oct 22–Nov 22, 2026' / 'Dec 22, 2026 – Jan 22, 2027' (대시가 빠져도 동작)."""
    seg = text[text.find(after):]
    m = re.search(r"([A-Z][a-z]{2}) (\d{1,2})(?:, (\d{4}))?\W+([A-Z][a-z]{2}) (\d{1,2}), (\d{4})", seg)
    if not m:
        return None
    m1, d1, y1, m2, d2, y2 = m.groups()
    y2 = int(y2)
    y1 = int(y1) if y1 else (y2 - 1 if MONTHS[m1] > MONTHS[m2] else y2)
    return dt.date(y1, MONTHS[m1], int(d1)), dt.date(y2, MONTHS[m2], int(d2))


def parse_invoice(path):
    t = pdf_text(path)
    inv = re.search(r"Invoice number\s+([A-Z0-9]+)\W+(\d{4})", t)
    issue = re.search(r"Date of issue\s+([A-Z][a-z]+ \d{1,2}, \d{4})", t)
    total = re.search(r"\bTotal\s+\$([\d,]+\.\d\d)", t)
    plan = re.search(r"Max plan\W+(\d+x)\s.*?\$([\d,]+\.\d\d)", t)
    if not (inv and issue and total and plan):
        die(f"인보이스 파싱 실패: {path}")
    credit = re.search(r"Unused time on Max plan\W+(\d+x).*?\$([\d,]+\.\d\d)", t)
    return {
        "invoice": f"{inv.group(1)}-{inv.group(2)}",
        "issue": long_date(issue.group(1)),
        "usd": money(total.group(1)),
        "plan": plan.group(1),
        "gross": money(plan.group(2)),
        "period": parse_period(t, plan.group(0)),
        "old_plan": credit.group(1) if credit else None,
        "credit": money(credit.group(2)) if credit else None,
    }


def parse_receipt(path):
    t = pdf_text(path)
    inv = re.search(r"Invoice number\s+([A-Z0-9]+)\W+(\d{4})", t)
    rno = re.search(r"Receipt number\s+(\d{4})\W+(\d{4})\W+(\d{4})", t)
    paid = re.search(r"Date paid\s+([A-Z][a-z]+ \d{1,2}, \d{4})", t)
    amt = re.search(r"Amount paid\s+\$([\d,]+\.\d\d)", t)
    card = re.search(r"(Visa|Mastercard|Amex)\W+(\d{4})", t)
    if not (inv and rno and paid and amt):
        die(f"영수증 파싱 실패: {path}")
    return {
        "invoice": f"{inv.group(1)}-{inv.group(2)}",
        "receipt": "-".join(rno.groups()),
        "paid": long_date(paid.group(1)),
        "usd": money(amt.group(1)),
        "card": f"{card.group(1)} {card.group(2)}" if card else None,
    }


# ---------------------------------------------------------------- 엑셀
def src_rects(xlsx):
    """템플릿 drawing 의 이미지 crop(srcRect)을 읽는다. openpyxl 은 저장할 때 이걸 버린다."""
    out = {}
    with zipfile.ZipFile(xlsx) as z:
        for name in z.namelist():
            if not re.match(r"xl/drawings/drawing\d+\.xml$", name):
                continue
            root = ET.fromstring(z.read(name))
            for anchor in root:
                frm = anchor.find(f"{XDR}from")
                rect = anchor.find(f".//{A}srcRect")
                if frm is None or rect is None:
                    continue
                key = (int(frm.find(f"{XDR}col").text), int(frm.find(f"{XDR}row").text))
                out[key] = {k: int(rect.get(k, 0)) for k in "ltrb"}
    return out


def bake_crop(img, rect):
    """srcRect(1/1000 %) 만큼 잘라낸 PNG 로 교체 -> 저장 후에도 로고가 원본과 똑같이 보인다."""
    pil = Image.open(io.BytesIO(img._data()))
    w, h = pil.size
    box = (round(w * rect["l"] / 1e5), round(h * rect["t"] / 1e5),
           round(w * (1 - rect["r"] / 1e5)), round(h * (1 - rect["b"] / 1e5)))
    buf = io.BytesIO()
    pil.crop(box).save(buf, "PNG")
    buf.seek(0)
    new = XLImage(buf)
    new.anchor = img.anchor
    return new


def ddmmyyyy(d):
    return d.strftime("%d/%m/%Y")


def build_xlsx(base, sig, out_path, inv, rec, thb, pr_no):
    rects = src_rects(base)
    wb = openpyxl.load_workbook(base)
    ws = wb.active

    ws["C6"] = NAME
    if not ws["C7"].value:
        ws["C7"] = "JKT"
    ws["H6"] = "=TODAY()"

    desc = f"ค่าบริการ AI รายเดือน Claude Max plan {inv['plan']} (Anthropic)"
    if inv["old_plan"] and inv["period"]:
        s, e = inv["period"]
        desc += f" อัปเกรดจาก {inv['old_plan']} รอบ {ddmmyyyy(s)}-{ddmmyyyy(e)}"

    remark = f"Invoice {inv['invoice']} USD {inv['usd']:,.2f}"
    if inv["credit"]:
        # 'Max 20x 200.00' 은 폰트 합자로 '20× 200.00' 처럼 찍혀서 숫자 사이에 태국어를 둔다
        remark += (f" (Max {inv['plan']} ราคา {inv['gross']:,.2f} หักเครดิต Max {inv['old_plan']}"
                   f" ที่เหลือ {inv['credit']:,.2f})")
    if thb is not None:
        remark += f" = THB {thb:,.2f}"
    remark += f" ตามยอดเรียกเก็บบัตร {ddmmyyyy(rec['paid'])}"
    if pr_no:
        remark += f" / อ้างอิง {pr_no}"

    ws["B10"] = 1
    ws["C10"] = desc
    ws["D10"] = 1
    ws["E10"] = "เดือน"
    ws["F10"] = thb  # None -> 빈칸 (환율 추정 금지)
    ws["G10"] = "=D10*F10"
    ws["H10"] = remark
    for ref in ("C10", "H10"):
        al = copy.copy(ws[ref].alignment)
        ws[ref].alignment = Alignment(horizontal="left", vertical=al.vertical or "center",
                                      wrap_text=True)
    f = copy.copy(ws["H10"].font)
    f.size = 11
    ws["H10"].font = f
    ws.row_dimensions[10].height = 72
    for ref in ("G18", "G19"):  # Vat 7% / WHT 비움
        ws[ref] = None

    # 이미지: 상단 로고만 남기고(crop 반영) 서명은 새로 앵커
    keep = []
    for img in ws._images:
        frm = img.anchor._from
        if frm.row < 6:
            rect = rects.get((frm.col, frm.row))
            keep.append(bake_crop(img, rect) if rect else img)
    ws._images = keep

    s = XLImage(sig)
    w = round(s.width * SIG_H_PX / s.height)
    s.width, s.height = w, SIG_H_PX
    s.anchor = OneCellAnchor(
        _from=AnchorMarker(col=2, colOff=pixels_to_EMU(40), row=21, rowOff=pixels_to_EMU(4)),
        ext=XDRPositiveSize2D(pixels_to_EMU(w), pixels_to_EMU(SIG_H_PX)))
    ws.add_image(s)
    wb.save(out_path)


# ---------------------------------------------------------------- PDF
def xlsx_to_pdf(xlsx, pdf_out):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        die("LibreOffice(soffice) 가 없습니다.")
    with tempfile.TemporaryDirectory() as td:
        src = os.path.join(td, "claim.xlsx")
        shutil.copy(xlsx, src)
        subprocess.run([soffice, f"-env:UserInstallation=file://{td}/profile", "--headless",
                        "--convert-to", "pdf", "--outdir", td, src],
                       check=True, capture_output=True)
        shutil.move(os.path.join(td, "claim.pdf"), pdf_out)


def sign_pr(pr, out_pdf):
    """PR 의 Review 칸 두 곳(요청부서·자재부서)에만 서명. jkt-sign 스크립트 재사용."""
    cands = sorted(glob.glob(os.path.expanduser("~/.claude/skills/**/jkt-sign/scripts/sign.py"),
                             recursive=True))
    if not cands:
        die("jkt-sign 스킬(sign.py)을 찾지 못했습니다. --pr-signed 로 서명본을 주세요.")
    subprocess.run([sys.executable, cands[0], "--pdf", pr, "--out", out_pdf,
                    "--anchor", "Review", "--occurrence", "0", "--height", "16"], check=True)


def mask_sms(src, boxes, out_png):
    im = Image.open(src).convert("RGB")
    from PIL import ImageDraw
    d = ImageDraw.Draw(im)
    for b in boxes:
        d.rectangle(tuple(int(v) for v in b.split(",")), fill=MASK_RGB)  # 글자 없이 회색 블록만
    im.save(out_png)
    return out_png


def image_pdf(png):
    buf = io.BytesIO()
    Image.open(png).convert("RGB").save(buf, "PDF", resolution=150)
    buf.seek(0)
    return buf


# ---------------------------------------------------------------- 검증
def verify(xlsx, claim_pdf, bundle, thb, verify_png):
    problems = []
    wb = openpyxl.load_workbook(xlsx)
    ws = wb.active
    n_img = len(ws._images)
    if n_img != 2:
        problems.append(f"이미지 {n_img}개 (로고+서명 2개여야 함)")
    if ws["G20"].value != "=SUM(G10:G19)":
        problems.append(f"G20 수식 변경됨: {ws['G20'].value}")

    import pdfplumber
    with pdfplumber.open(claim_pdf) as pdf:
        pages = len(pdf.pages)
        text = pdf.pages[0].extract_text() or ""
    if pages != 1:
        problems.append(f"ใบเบิก PDF 가 {pages}쪽 (1쪽이어야 함)")
    errs = re.findall(r"#(?:VALUE|REF|NAME|DIV/0|N/A|NUM|NULL)[!?]?", text)
    if errs:
        problems.append(f"수식 오류 {errs}")
    gt = re.search(r"Grand Total\s+([\d,]+\.\d\d|-)", text)
    gt_val = gt.group(1) if gt else "?"
    if thb is not None and gt_val != f"{thb:,.2f}":
        problems.append(f"Grand Total {gt_val} ≠ THB {thb:,.2f}")

    try:
        import pypdfium2 as pdfium
        pdfium.PdfDocument(claim_pdf)[0].render(scale=2).to_pil().save(verify_png)
    except ImportError:
        subprocess.run(["pdftoppm", "-png", "-r", "144", "-singlefile", claim_pdf,
                        verify_png[:-4]], check=True)
    return problems, gt_val, len(PdfReader(bundle).pages)


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--root", default=os.path.join(REPO, "anthropic-claim"))
    ap.add_argument("--month", required=True, help="inbox 하위 폴더 YYMM (예: 2610)")
    ap.add_argument("--base", help=f"ใบเบิก 원본 양식 (기본 template/{TEMPLATE_NAME})")
    ap.add_argument("--thb", type=float, help="바트 금액: PR Price -> KBank 문자 순. 없으면 생략")
    ap.add_argument("--pr", help="사내 PR PDF (서명 전 원본)")
    ap.add_argument("--pr-signed", help="이미 Review 칸 서명된 PR PDF")
    ap.add_argument("--pr-no", help="PR 번호 (기본: PR 파일명에서)")
    ap.add_argument("--sms", help="KBank 문자 캡처")
    ap.add_argument("--mask", action="append", default=[], help="가릴 영역 x0,y0,x1,y1 (px)")
    ap.add_argument("--date", help="작성일 YYMMDD (기본 오늘)")
    a = ap.parse_args()

    root = os.path.abspath(a.root)
    inbox = os.path.join(root, "inbox", a.month)
    out = os.path.join(root, "out")
    os.makedirs(out, exist_ok=True)

    base = a.base or os.path.join(root, "template", TEMPLATE_NAME)
    if not a.base and not os.path.exists(base):
        # 원본 양식이 아직 없으면 지난달 완성본을 바탕으로 쓴다 (10행·서명은 어차피 새로 씀)
        prev = sorted(p for p in glob.glob(os.path.join(out, "*_ใบเบิก_Anthropic_*.xlsx"))
                      if "_DRAFT" not in p)
        if prev:
            base = prev[-1]
            print(f"[알림] template/{TEMPLATE_NAME} 없음 -> 지난 완성본 사용: {os.path.basename(base)}")
    sig = os.path.join(root, "template", "sig_chin.png")
    for p, what in ((base, "ใบเบิก 양식"), (sig, "서명 PNG")):
        if not os.path.exists(p):
            die(f"{what} 없음: {p}")

    invs = sorted(glob.glob(os.path.join(inbox, "Invoice-*.pdf")))
    recs = sorted(glob.glob(os.path.join(inbox, "Receipt-*.pdf")))
    if len(invs) != 1 or len(recs) != 1:
        die(f"{inbox} 에 Invoice/Receipt 가 각 1개씩 있어야 합니다 (Invoice {len(invs)}, Receipt {len(recs)}).")
    inv, rec = parse_invoice(invs[0]), parse_receipt(recs[0])
    if inv["invoice"] != rec["invoice"] or abs(inv["usd"] - rec["usd"]) > 0.005:
        die(f"인보이스/영수증 불일치: {inv['invoice']} {inv['usd']} vs {rec['invoice']} {rec['usd']}")

    pr = a.pr_signed or a.pr
    pr_no = a.pr_no or (re.search(r"PR-\d{8}", os.path.basename(pr)).group(0)
                        if pr and re.search(r"PR-\d{8}", os.path.basename(pr)) else None)
    if a.pr and a.thb is not None:
        if f"{a.thb:,.2f}" not in pdf_text(a.pr) and f"{a.thb:.2f}" not in pdf_text(a.pr):
            print(f"[주의] PR 본문에서 {a.thb:,.2f} 를 찾지 못했습니다. Price 칸을 다시 확인하세요.")

    ymd = a.date or dt.date.today().strftime("%y%m%d")
    stem = f"{ymd}_ใบเบิก_Anthropic_{NAME}"
    draft = "_DRAFT" if a.thb is None else ""
    xlsx = os.path.join(out, f"{stem}{draft}.xlsx")
    claim_pdf = os.path.join(out, f"{stem}{draft}.pdf")
    bundle = os.path.join(out, f"{stem}_전체{draft}.pdf")
    verify_png = os.path.join(out, f"{stem}{draft}_p1_check.png")

    build_xlsx(base, sig, xlsx, inv, rec, a.thb, pr_no)
    xlsx_to_pdf(xlsx, claim_pdf)

    parts = [("ใบเบิก", claim_pdf), ("Invoice", invs[0]), ("Receipt", recs[0])]
    if a.pr_signed:
        parts.append(("PR(서명본)", a.pr_signed))
    elif a.pr:
        signed = os.path.join(out, os.path.basename(a.pr)[:-4] + "_signed.pdf")
        sign_pr(a.pr, signed)
        parts.append(("PR(서명)", signed))
    if a.sms:
        sms_png = os.path.join(out, "sms_masked.png") if a.mask else a.sms
        if a.mask:
            mask_sms(a.sms, a.mask, sms_png)
        parts.append(("문자", sms_png))

    w = PdfWriter()
    for _, p in parts:
        w.append(image_pdf(p) if p.lower().endswith((".png", ".jpg", ".jpeg")) else p)
    with open(bundle, "wb") as fh:
        w.write(fh)

    problems, gt, n_pages = verify(xlsx, claim_pdf, bundle, a.thb, verify_png)

    print("\n=== 정산 결과 ===")
    print(f"Invoice   {inv['invoice']}  발행 {inv['issue']}  Max {inv['plan']}"
          + (f" (업그레이드: {inv['gross']:,.2f} − {inv['old_plan']} 크레딧 {inv['credit']:,.2f})"
             if inv["credit"] else ""))
    print(f"Receipt   {rec['receipt']}  결제 {rec['paid']}  {rec['card'] or ''}")
    print(f"USD       {inv['usd']:,.2f}")
    print(f"THB       {f'{a.thb:,.2f}' if a.thb is not None else '(비움 — 금액 근거 필요)'}"
          f"   Grand Total(PDF) {gt}")
    print(f"PR        {pr_no or '(없음)'}")
    print("첨부 순서 " + " → ".join(k for k, _ in parts) + f"  ({n_pages}쪽)")
    print(f"엑셀      {xlsx}\n합본      {bundle}\n확인용    {verify_png}")
    print("검증      " + ("OK" if not problems else "; ".join(problems)))


if __name__ == "__main__":
    main()
